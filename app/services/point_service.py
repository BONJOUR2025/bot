"""Кабинет точки на рабочем ПК салона — данные экрана смены.

Точка — салон из «Салонов», её склады приёма — salon.sclad_ids (для Бестужевской
добавляем и склад приёмки в том же здании). Всё только чтение Агбиса; своё —
отметки звонков (point_device_repository.PointCallRepository).

Три списка, на которые смена смотрит весь день:
- «Выдача» — заказы точки в статусе «Исполненный»: изделие готово, клиент ещё
  не забрал. Позвонить надо тем, по кому ещё нет отметки звонка.
- «Сроки» — не готовые заказы со сроком сегодня или раньше: клиента нужно
  предупредить до того, как он приедет.
- «Долго лежат» — готовы больше STALE_DAYS дней назад и до сих пор не выданы.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from typing import Any

STALE_DAYS = 14
DUE_WINDOW_DAYS = 60
READY = 4
OPEN_STATUSES = (1, 2, 3)


def point_sclads(salon) -> list[int]:
    from app.services.workshop_order_service import BESTUZHEVSKAYA_SCLADS

    ids = {int(x) for x in (getattr(salon, "sclad_ids", None) or [])}
    if ids & BESTUZHEVSKAYA_SCLADS:
        ids |= BESTUZHEVSKAYA_SCLADS
    return sorted(ids)


def _money(v) -> float:
    try:
        return round(float(v or 0), 2)
    except Exception:
        return 0.0


def orders(salon) -> dict[str, Any]:
    from app.services.firebird_service import _connect
    from app.services.workshop_order_service import _iso, _text

    sclads = point_sclads(salon)
    if not sclads:
        return {"ready": [], "due": [], "sclads": []}
    now = datetime.now()
    today_end = datetime.combine(date.today(), datetime.max.time())
    con = _connect()
    try:
        cur = con.cursor()
        marks = ",".join("?" * len(sclads))
        cur.execute(
            f"""
            SELECT dor.id, d.doc_num, d.doc_date, dor.status_id, dor.date_out, dor.date_complete,
                   dor.kredit, dor.debet, dor.current_sclad_id, dor.fast_execute, d.contragent_id,
                   c.name, c.teleph_cell, dor.contact_tel1
            FROM docs_order dor
                JOIN docs d ON d.doc_id = dor.doc_id
                LEFT JOIN contragents c ON c.contr_id = d.contragent_id
            WHERE dor.sclad_kredit_id IN ({marks})
              AND (dor.status_id = {READY}
                   OR (dor.status_id IN ({",".join(map(str, OPEN_STATUSES))}) AND dor.date_out <= ?))
              AND d.doc_date > DATEADD(-400 DAY TO CURRENT_DATE)
            """, (*sclads, today_end))
        rows = cur.fetchall()
        ids = [r[0] for r in rows]

        # Что изделия — по одной строке-изделию на заказ, чтобы в списке было
        # «Туфли, Сумка», а не номер без смысла.
        items: dict[int, list[str]] = {}
        sms: dict[int, dict[str, Any]] = {}
        for start in range(0, len(ids), 200):
            chunk = ids[start:start + 200]
            m = ",".join("?" * len(chunk))
            cur.execute(
                f"SELECT dos.doc_order_id, t.name FROM doc_order_services dos "
                f"LEFT JOIN tovars_tbl t ON t.tovar_id = dos.tovar_id "
                f"WHERE dos.doc_order_id IN ({m}) AND dos.parent_dos_id IS NULL ORDER BY dos.id", chunk)
            for oid, name in cur.fetchall():
                n = _text(name)
                if n:
                    items.setdefault(oid, []).append(n)
        doc_ids = {}
        if ids:
            for start in range(0, len(ids), 200):
                chunk = ids[start:start + 200]
                m = ",".join("?" * len(chunk))
                cur.execute(f"SELECT id, doc_id FROM docs_order WHERE id IN ({m})", chunk)
                doc_ids.update({doc: oid for oid, doc in cur.fetchall()})
            docs = list(doc_ids)
            for start in range(0, len(docs), 200):
                chunk = docs[start:start + 200]
                m = ",".join("?" * len(chunk))
                cur.execute(
                    f"SELECT doc_id, dt_added, dt_deliver FROM smses WHERE doc_id IN ({m}) ORDER BY dt_added", chunk)
                for doc, added, delivered in cur.fetchall():
                    oid = doc_ids.get(doc)
                    if oid is not None:
                        sms[oid] = {"sent": _iso(added), "delivered": _iso(delivered)}
    finally:
        con.close()

    from app.data.point_device_repository import PointCallRepository
    from app.services.workshop_order_service import sclad_label

    calls = PointCallRepository().last_for(ids)
    ready, due = [], []
    for (oid, num, ddate, status, date_out, completed, kredit, debet, cur_sclad, fast, cid,
         cname, cphone, tel1) in rows:
        due_dt = date_out if isinstance(date_out, datetime) and date_out.year > 2000 else None
        row = {
            "order_id": oid,
            "doc_num": _text(num),
            "accepted": _iso(ddate),
            "items": items.get(oid, []),
            "client": _text(cname),
            "phone": _text(cphone) or _text(tel1),
            "to_pay": max(0.0, _money(kredit) - _money(debet)),
            "total": _money(kredit),
            "urgent": bool(fast),
            "due": _iso(due_dt),
            "location": sclad_label(cur_sclad),
            "sms": sms.get(oid),
            "call": calls.get(oid),
        }
        if status == READY:
            done = completed if isinstance(completed, datetime) else None
            row["ready_since"] = _iso(done)
            row["ready_days"] = (now - done).days if done else None
            row["stale"] = bool(done and (now - done) >= timedelta(days=STALE_DAYS))
            ready.append(row)
        else:
            row["overdue_days"] = (now.date() - due_dt.date()).days if due_dt else None
            due.append(row)
    ready.sort(key=lambda r: (r["call"] is not None, r["ready_since"] or ""))
    # Просрочка старше DUE_WINDOW_DAYS — это уже не «предупредить клиента до
    # приезда», а забытые заказы (часто индивидуальный пошив): отдельным счётчиком.
    old_due = [r for r in due if (r["overdue_days"] or 0) > DUE_WINDOW_DAYS]
    due = [r for r in due if (r["overdue_days"] or 0) <= DUE_WINDOW_DAYS]
    due.sort(key=lambda r: r["due"] or "")
    return {"ready": ready, "due": due, "old_due_count": len(old_due), "sclads": sclads}


def shift(salon) -> dict[str, Any]:
    """Кто сегодня отметился на точке (та же отметка, что приходит из бота)."""
    from app.data.employee_repository import EmployeeRepository
    from app.data.shift_checkin_repository import get_shift_checkin_repository

    today = date.today().isoformat()
    rows = get_shift_checkin_repository().list(date_from=today, date_to=today, salon_id=salon.id)
    emps = EmployeeRepository()
    people, seen = [], set()
    for r in sorted(rows, key=lambda x: x.get("sent_at", "")):
        eid = str(r.get("employee_id") or "")
        if eid in seen:
            continue
        seen.add(eid)
        e = emps.get_employee(eid)
        people.append({"name": (e.name if e else "") or r.get("employee_name") or "Сотрудник",
                       "at": r.get("sent_at")})
    return {"opened": bool(people), "people": people}


# ── приём сегодня и перемещения точка ↔ цех ─────────────────────────
def accepted_today(salon) -> list[dict[str, Any]]:
    """Заказы, принятые на точке сегодня: что именно и на кого."""
    from app.services.firebird_service import _connect
    from app.services.workshop_order_service import STATUS_NAMES, _iso, _text

    sclads = point_sclads(salon)
    if not sclads:
        return []
    con = _connect()
    try:
        cur = con.cursor()
        m = ",".join("?" * len(sclads))
        cur.execute(
            f"""
            SELECT dor.id, d.doc_num, d.doc_date, d.doc_time, dor.status_id, dor.date_out, dor.fast_execute,
                   c.name
            FROM docs_order dor JOIN docs d ON d.doc_id = dor.doc_id
                LEFT JOIN contragents c ON c.contr_id = d.contragent_id
            WHERE dor.sclad_kredit_id IN ({m}) AND d.doc_date = CURRENT_DATE AND dor.status_id <> 7
            ORDER BY d.doc_time
            """, sclads)
        rows = cur.fetchall()
        items: dict[int, list[str]] = {}
        ids = [r[0] for r in rows]
        if ids:
            mm = ",".join("?" * len(ids))
            cur.execute(
                f"SELECT dos.doc_order_id, t.name FROM doc_order_services dos LEFT JOIN tovars_tbl t ON t.tovar_id = dos.tovar_id "
                f"WHERE dos.doc_order_id IN ({mm}) AND dos.parent_dos_id IS NULL ORDER BY dos.id", ids)
            for oid, name in cur.fetchall():
                if _text(name):
                    items.setdefault(oid, []).append(_text(name))
    finally:
        con.close()
    out = []
    for oid, num, ddate, dtime, status, date_out, fast, cname in rows:
        due = date_out if isinstance(date_out, datetime) and date_out.year > 2000 else None
        out.append({"order_id": oid, "doc_num": _text(num), "time": str(dtime)[:5] if dtime else "",
                    "status": STATUS_NAMES.get(status, "—"), "due": _iso(due), "urgent": bool(fast),
                    "client": _text(cname), "items": items.get(oid, [])})
    return out


def logistics(salon) -> dict[str, Any]:
    """Накладные точки за последние две недели: что едет к нам и что от нас.

    Статусы Агбиса: 1 — готова к отгрузке, 2 — в пути, 3 — принята,
    4 — принята не полностью. «Не полностью» — повод проверить полку."""
    from app.services.firebird_service import _connect
    from app.services.workshop_order_service import IN_WAY_STATUS, _iso, _text, sclad_label

    sclads = point_sclads(salon)
    if not sclads:
        return {"incoming": [], "outgoing": []}
    con = _connect()
    try:
        cur = con.cursor()
        m = ",".join("?" * len(sclads))
        cur.execute(
            f"""
            SELECT w.id, w.doc_num, w.doc_date, w.from_sclad_id, w.to_sclad_id, w.diw_status_id,
                   (SELECT COUNT(*) FROM docs_in_way_servs s WHERE s.doc_in_way_id = w.id)
            FROM docs_in_way w
            WHERE w.doc_date > DATEADD(-14 DAY TO CURRENT_DATE)
              AND (w.to_sclad_id IN ({m}) OR w.from_sclad_id IN ({m}))
            ORDER BY w.doc_date DESC, w.id DESC
            """, (*sclads, *sclads))
        rows = cur.fetchall()
    finally:
        con.close()
    incoming, outgoing = [], []
    for wid, num, ddate, frm, to, st, cnt in rows:
        # Перемещения внутри одной точки (приёмка ↔ цех в одном здании) — не логистика.
        if frm in sclads and to in sclads:
            continue
        row = {"id": wid, "doc_num": _text(num), "date": _iso(ddate), "from": sclad_label(frm),
               "to": sclad_label(to), "status": IN_WAY_STATUS.get(st, "—"), "status_id": st, "items": int(cnt or 0)}
        (incoming if to in sclads else outgoing).append(row)
    return {"incoming": incoming, "outgoing": outgoing}


# ── график точки ──────────────────────────────────────────────────────
_WEEK_CACHE: dict[str, tuple[float, list]] = {}
_WEEK_TTL = 600


def _read_week(code: str, days: int) -> list[dict[str, Any]]:
    """Одно открытие Excel-графика на всю неделю (read_only — в разы быстрее
    полной загрузки). Раньше по дню открывали книгу заново: 7 × 4 секунды."""
    import os
    from datetime import date as _date

    from openpyxl import load_workbook

    from app.config import EXCEL_FILE
    from app.core.constants import MONTHS_RU

    dates = [_date.today() + timedelta(days=i) for i in range(days)]
    who: dict[str, str] = {}
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE, data_only=True, read_only=True)
        try:
            for (y, m) in sorted({(d.year, d.month) for d in dates}):
                name = MONTHS_RU[m - 1]
                title = next((t for t in wb.sheetnames if t.upper() == name.upper()), None)                     or next((t for t in wb.sheetnames if t.upper().startswith(name.upper())), None)
                if not title:
                    continue
                rows = list(wb[title].iter_rows(min_row=1, max_row=40, max_col=40, values_only=True))
                if len(rows) < 3:
                    continue
                # Столбец дня — по номеру дня в первой или второй строке.
                col_of: dict[int, int] = {}
                for c in range(len(rows[0])):
                    for hdr in (rows[0][c], rows[1][c] if len(rows) > 1 else None):
                        try:
                            n = int(str(hdr).strip())
                        except (TypeError, ValueError):
                            continue
                        if 1 <= n <= 31 and n not in col_of:
                            col_of[n] = c
                for d in dates:
                    if (d.year, d.month) != (y, m) or d.day not in col_of:
                        continue
                    c = col_of[d.day]
                    names = [str(r[0]).strip() for r in rows[2:]
                             if r and r[0] and c < len(r) and str(r[c] or "").strip() == code]
                    who[d.isoformat()] = ", ".join(names)
        finally:
            wb.close()
    return [{"date": d.isoformat(), "employee": who.get(d.isoformat(), "")} for d in dates]


async def week_schedule(salon, days: int = 7) -> list[dict[str, Any]]:
    """Кто работает на точке ближайшие дни — из того же Excel-графика, что
    видит бот. Точка в графике — код салона (у Бестужевской «Ц»)."""
    import asyncio
    import time

    code = (salon.code or "").strip()
    hit = _WEEK_CACHE.get(code)
    if hit and time.time() - hit[0] < _WEEK_TTL:
        return hit[1]
    data = await asyncio.to_thread(_read_week, code, days)
    _WEEK_CACHE[code] = (time.time(), data)
    return data


# ── база знаний ───────────────────────────────────────────────────────
# Что администратору на стойке знать не нужно: условия оплаты сотрудников.
KB_HIDDEN_CATEGORIES = {"Авто-обучение"}


def kb_documents() -> list[dict[str, Any]]:
    from app.db.session import SessionLocal
    from app.models.knowledge import KnowledgeDocument

    db = SessionLocal()
    try:
        docs = db.query(KnowledgeDocument).order_by(KnowledgeDocument.order_idx, KnowledgeDocument.id).all()
        return [{"id": d.id, "title": d.title, "category": d.category, "content": d.content}
                for d in docs
                if (d.content or "").strip() and (d.category or "") not in KB_HIDDEN_CATEGORIES]
    finally:
        db.close()


def kb_ask(question: str, salon_name: str) -> str:
    """Ответ по базе знаний — как «📚 База знаний» в боте, та же модель."""
    from app.handlers.knowledge_base import _kb_model
    from app.services.config_service import ConfigService
    from app.services.llm_client import chat, get_client

    cfg = ConfigService().load()
    if not get_client(cfg):
        raise RuntimeError("Помощник не настроен: нет ключа нейросети.")
    kb = "\n\n".join(f"=== {d['title']} ({d['category']}) ===\n{d['content'].strip()}" for d in kb_documents())
    system = f"""Ты помощник администратора салона BONJOUR (точка «{salon_name}»). Отвечай на вопросы строго по базе знаний ниже: прайс, методички, регламенты, инструкции по оформлению.

База знаний компании:
{kb}

Правила:
1. Отвечай только по базе знаний. Не придумывай цены, сроки и услуги.
2. Если ответа в базе нет — скажи: «В базе знаний этого нет. Уточните у руководителя или мастера цеха.»
3. Если спрашивают цену — назови услугу так, как она названа в прайсе, и цену.
4. Пиши кратко, по делу, без форматирования (никаких **, *, #). На «вы»."""
    reply = chat(cfg, [{"role": "user", "content": question}], system=system, max_tokens=700,
                 model=_kb_model(cfg), employee_id=f"point:{salon_name}", employee_name=f"Точка {salon_name}",
                 feature="point_knowledge_base") or ""
    return reply.replace("**", "").replace("__", "").strip() or "Помощник не ответил, попробуйте ещё раз."
